#!/usr/bin/env python3
"""
카페24 쇼핑몰 상품 이미지 깨짐 검사 (기본: vivacious.kr)
================================================================
쇼핑몰의 모든 상품 상세 페이지를 찾아 상품 이미지(대표·추가·상세 이미지)를
하나하나 내려받아 검사하고, 문제가 있는 상품을 엑셀로 정리합니다.

검사 항목
  [깨짐]  HTTP 오류(404 등)로 이미지 누락 / 0바이트 / 손상되어 열리지 않음
          파일이 중간에 잘림(업로드·전송 중 끊김) / 이미지 대신 오류 페이지 응답
  [주의]  하단이 회색 띠로 채워진 '잘린 채 저장된' 이미지 의심
          상품 페이지에 대표 이미지 또는 상세 이미지가 하나도 없음
          비정상적으로 작은 이미지(1~2px)

설치 (최초 1회)
    pip install requests beautifulsoup4 pillow openpyxl

실행
    python shop_image_check.py
    python shop_image_check.py https://다른카페24몰.com     # 다른 카페24 몰
    python shop_image_check.py --limit 50                  # 상품 50개만 시험 실행

결과
    vivacious_이미지검사_YYYYMMDD_HHMM.xlsx
      - '문제 이미지' 시트 : 상품명, 상품 링크, 문제 유형, 이미지 URL
      - '상품별 요약' 시트 : 전체 상품의 이미지 수 / 문제 수
"""

import argparse
import io
import re
import sys
import threading
import time
from collections import defaultdict
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime
from urllib.parse import urljoin, urlparse, parse_qs, urlencode, unquote, quote

import requests
from bs4 import BeautifulSoup
from PIL import Image, ImageFile, ImageStat

for _s in (sys.stdout, sys.stderr):  # 윈도우 콘솔(cp949)에서 특수문자 출력 오류 방지
    try:
        _s.reconfigure(errors="replace")
    except Exception:
        pass

ImageFile.LOAD_TRUNCATED_IMAGES = False
Image.MAX_IMAGE_PIXELS = None
try:
    import pillow_avif  # noqa: F401  (AVIF 지원, 선택)
except ImportError:
    pass

UA = ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
      "(KHTML, like Gecko) Chrome/130.0 Safari/537.36")

PRODUCT_RE = re.compile(r"^/product/([^/]+)/(\d+)(?:/|$)")
LIST_PATHS = ("/product/list.html", "/category/")
SKIP_PATHS = ("/myshop", "/order", "/member", "/board", "/product/search", "/exec/",
              "/product/compare", "/product/recent", "/shopinfo", "/link/", "/attend", "/coupon")
LAZY_ATTRS = ("src", "data-src", "ec-data-src", "data-original", "data-lazy", "data-lazy-src", "data-img")
BG_URL_RE = re.compile(r"url\(\s*['\"]?([^'\")]+)['\"]?\s*\)", re.I)
DETAIL_SELECTORS = ("#prdDetail", ".xans-product-additional", "#prdInfo", ".detail_cont",
                    ".prdDetail", "#detail", ".cont_detail")
# 쇼핑몰 공통 장식 이미지(아이콘, 버튼, 배너)는 상품 이미지에서 제외
DECOR_RE = re.compile(r"(/0dzimg/|/24d/|/web/upload/(icon|btn|ico)|/skin|/_wg/|/layout/|ico_|btn_|icon_|"
                      r"/img/(common|btn|icon)|\.svg(\?|$)|/web/upload/category/)", re.I)

_local = threading.local()


def S():
    s = getattr(_local, "s", None)
    if s is None:
        s = requests.Session()
        s.headers.update({"User-Agent": UA, "Accept-Language": "ko-KR,ko;q=0.9"})
        _local.s = s
    return s


# ------------------------------------------------------------------ URL 정리
def canon_product(url):
    p = urlparse(url)
    m = PRODUCT_RE.match(p.path)
    if not m:
        return None, None
    slug = quote(unquote(m.group(1)), safe="")  # 한글 상품 주소를 %인코딩으로 통일
    return f"{p.scheme}://{p.netloc}/product/{slug}/{m.group(2)}/", m.group(2)


def canon_list(url):
    """카테고리 목록 URL은 cate_no, page 파라미터만 남긴다."""
    p = urlparse(url)
    q = parse_qs(p.query)
    keep = {k: q[k][0] for k in ("cate_no", "page") if k in q}
    path = p.path
    return f"{p.scheme}://{p.netloc}{path}" + (f"?{urlencode(sorted(keep.items()))}" if keep else "")


def is_list(path):
    return any(path.startswith(x) for x in LIST_PATHS)


# ------------------------------------------------------------------ 상품 URL 수집
def from_sitemap(base):
    found = {}
    todo, seen = [urljoin(base, "/sitemap.xml")], set()
    while todo:
        sm = todo.pop()
        if sm in seen or len(seen) > 50:
            continue
        seen.add(sm)
        try:
            r = S().get(sm, timeout=20)
            if r.status_code != 200:
                continue
        except Exception:
            continue
        for loc in re.findall(r"<loc>\s*([^<]+?)\s*</loc>", r.text):
            loc = loc.replace("&amp;", "&")
            if loc.lower().endswith((".xml", ".xml.gz")):
                todo.append(loc)
            else:
                cu, no = canon_product(loc)
                if cu and no not in found:
                    found[no] = cu
    return found


def discover_products(base, workers, max_list_pages, delay):
    domain = urlparse(base).netloc
    products = from_sitemap(base)
    print(f"  사이트맵에서 상품 {len(products)}개 확인")

    queue, seen, visited = [canon_list(base)], {canon_list(base)}, 0
    lock = threading.Lock()

    def visit(url):
        try:
            r = S().get(url, timeout=20)
            if r.status_code != 200 or "html" not in r.headers.get("Content-Type", ""):
                return []
            r.encoding = r.encoding if r.encoding and r.encoding.lower() != "iso-8859-1" else "utf-8"
            soup = BeautifulSoup(r.text, "html.parser")
        except Exception:
            return []
        nxt = []
        for a in soup.find_all("a", href=True):
            h = a["href"].strip()
            if h.startswith(("javascript:", "mailto:", "tel:", "#")):
                continue
            u = urljoin(r.url, h)
            p = urlparse(u)
            if p.netloc != domain:
                continue
            cu, no = canon_product(u)
            if cu:
                with lock:
                    products.setdefault(no, cu)
                continue
            if any(p.path.startswith(x) for x in SKIP_PATHS):
                continue
            if is_list(p.path) or p.path in ("", "/", "/index.html"):
                nxt.append(canon_list(u))
        if delay:
            time.sleep(delay)
        return nxt

    with ThreadPoolExecutor(max_workers=workers) as ex:
        while queue and visited < max_list_pages:
            batch, queue = queue[:max_list_pages - visited], []
            visited += len(batch)
            for res in ex.map(visit, batch):
                for u in res:
                    if u not in seen:
                        seen.add(u)
                        queue.append(u)
            print(f"  목록 페이지 {visited}개 탐색 · 누적 상품 {len(products)}개")

    # '더보기' 버튼 방식 목록 대응: 카테고리마다 page=1,2,3... 를 새 상품이 안 나올 때까지 넘김
    cats = set()
    for u in seen:
        m = re.search(r"cate_no=(\d+)", u) or re.search(r"/category/[^/]+/(\d+)/?", u)
        if m:
            cats.add(m.group(1))
    scheme = urlparse(base).scheme

    def page_ids(cate, page):
        url = f"{scheme}://{domain}/product/list.html?cate_no={cate}&page={page}"
        try:
            r = S().get(url, timeout=20)
            soup = BeautifulSoup(r.text, "html.parser")
        except Exception:
            return {}
        out = {}
        for a in soup.find_all("a", href=True):
            cu, no = canon_product(urljoin(r.url, a["href"]))
            if cu:
                out[no] = cu
        return out

    def walk(cate):
        got, page = {}, 1
        while page <= 500:
            ids = page_ids(cate, page)
            new = {k: v for k, v in ids.items() if k not in got}
            if not new:
                break
            got.update(new)
            page += 1
            if delay:
                time.sleep(delay)
        return cate, page - 1, got

    if cats:
        print(f"  카테고리 {len(cats)}개 페이지 넘김 탐색 중...")
        with ThreadPoolExecutor(max_workers=workers) as ex:
            for cate, pages, got in ex.map(walk, sorted(cats, key=int)):
                for no, cu in got.items():
                    products.setdefault(no, cu)
                print(f"    카테고리 {cate}: {pages}페이지 · 상품 {len(got)}개 (누적 {len(products)}개)")
    return products


# ------------------------------------------------------------------ 상품 페이지 분석
def parse_product(url, base_domain):
    slug_name = unquote(urlparse(url).path.split("/")[2]).replace("-", " ")
    info = {"url": url, "name": slug_name, "images": {}, "notes": [], "error": ""}
    try:
        r = S().get(url, timeout=25)
        if r.status_code != 200:
            info["error"] = f"상품 페이지 HTTP {r.status_code}"
            return info
        r.encoding = r.encoding if r.encoding and r.encoding.lower() != "iso-8859-1" else "utf-8"
        soup = BeautifulSoup(r.text, "html.parser")
    except Exception as e:
        info["error"] = f"상품 페이지 접속 실패: {str(e)[:60]}"
        return info

    og = soup.find("meta", attrs={"property": "og:title"})
    name = (og.get("content") if og else "") or ""
    if not name:
        h = soup.select_one(".headingArea h2, .infoArea h2, .prd_name, h2")
        name = h.get_text(strip=True) if h else ""
    if not name and soup.title:
        name = soup.title.get_text(strip=True)
    info["name"] = name.strip() or slug_name

    imgs = info["images"]  # url -> 위치(대표/추가/상세)

    def add(u, where):
        if not u:
            return
        u = u.strip()
        if not u or u.startswith(("data:", "javascript:")):
            return
        if u.startswith("//"):
            u = "https:" + u
        u = urljoin(r.url, u).split("#")[0]
        if DECOR_RE.search(u):
            return
        imgs.setdefault(u, where)

    def collect(node, where):
        for t in node.find_all(["img", "source"]):
            for a in LAZY_ATTRS:
                if t.name == "source" and a == "src":
                    continue
                v = t.get(a) or ""
                if where == "상세" and "/web/product/" in v:
                    continue  # 상세 영역 안의 관련/추천 상품 썸네일은 제외
                add(v, where)
            for a in ("srcset", "data-srcset"):
                if t.get(a):
                    for part in t[a].split(","):
                        add(part.strip().split(" ")[0], where)
        for t in node.find_all(style=True):
            for u in BG_URL_RE.findall(t["style"]):
                add(u, where)

    # 상세 영역
    detail_nodes = [n for sel in DETAIL_SELECTORS for n in soup.select(sel)]
    for n in detail_nodes:
        collect(n, "상세")
    detail_count = sum(1 for v in imgs.values() if v == "상세")

    # 대표/추가 이미지 (카페24: /web/product/big|medium|small|extra/)
    meta_img = soup.find("meta", attrs={"property": "og:image"})
    if meta_img:
        add(meta_img.get("content"), "대표")
    for t in soup.select(".keyImg img, .BigImage, .xans-product-image img, .ThumbImage, "
                         ".xans-product-addimage img, .detailArea img"):
        for a in LAZY_ATTRS:
            add(t.get(a), "대표/추가")
    for t in soup.find_all("img"):
        for a in LAZY_ATTRS:
            v = t.get(a) or ""
            if "/web/product/" in v and ("/big/" in v or "/extra/" in v):
                add(v, "대표/추가")
    main_count = sum(1 for v in imgs.values() if v in ("대표", "대표/추가"))

    if main_count == 0:
        info["notes"].append("대표 이미지 없음")
    if detail_nodes and detail_count == 0:
        info["notes"].append("상세 이미지 없음")
    elif not detail_nodes:
        info["notes"].append("상세 영역을 찾지 못함(구조 확인 필요)")
    return info


# ------------------------------------------------------------------ 이미지 검사
def gray_band_bottom(im):
    """하단이 단색 회색으로 채워졌는지 (잘린 이미지를 그대로 다시 저장한 경우)."""
    w, h = im.size
    if h < 60 or w < 20:
        return False
    small = im.convert("RGB").resize((64, min(h, 800)))
    sh = small.size[1]

    def row_gray(y):
        st = ImageStat.Stat(small.crop((0, y, 64, y + 1)))
        r, g, b = st.mean
        return max(st.stddev) < 3 and 100 <= r <= 150 and abs(r - g) < 5 and abs(g - b) < 5

    run = 0
    while run < sh and row_gray(sh - 1 - run):
        run += 1
    if run < max(4, sh * 0.04) or run >= sh * 0.9:
        return False
    # 잘린 이미지: 회색 띠 바로 위 줄이 가로 전체에 걸쳐 '회색이 아닌 내용'으로 끊겨 있음
    # (회색 배경 촬영 사진은 피사체 양옆에 같은 회색이 이어지므로 제외됨)
    top = sh - run
    gray = small.getpixel((32, sh - 1))
    for y in range(max(0, top - 3), top):
        row = [small.getpixel((x, y)) for x in range(64)]
        same = sum(1 for p in row if max(abs(p[i] - gray[i]) for i in range(3)) < 12)
        if same > 64 * 0.15:
            return False
    return True


def check_image(url, referer, timeout, retries):
    referer = quote(referer, safe=":/?&=%#")  # 헤더에는 한글을 넣을 수 없음
    url = quote(url, safe=":/?&=%#+,;@~")
    res = {"status": "OK", "problem": "", "http": "", "size": "", "dim": ""}
    data, r = None, None
    for attempt in range(retries + 1):
        try:
            r = S().get(url, timeout=timeout, headers={"Referer": referer}, allow_redirects=False)
            hops = 0
            while r.is_redirect and hops < 5:
                loc = urljoin(r.url, r.headers.get("Location", ""))
                if re.search(r"serviceExpire|servicestop|hostinfo\.cafe24", loc, re.I):
                    return {**res, "status": "BROKEN", "http": r.status_code,
                            "problem": "이미지 서버 만료 (카페24 서비스 중지 페이지로 이동)"}
                r = S().get(loc, timeout=timeout, headers={"Referer": referer}, allow_redirects=False)
                hops += 1
            data = r.content
            if r.status_code >= 500 and attempt < retries:
                time.sleep(1.5)
                continue
            break
        except requests.exceptions.Timeout:
            res["problem"] = "타임아웃"
        except Exception as e:
            res["problem"] = f"연결 실패: {str(e)[:60]}"
        time.sleep(1.5)
    if data is None:
        res["status"] = "BROKEN"
        return res

    res["http"], res["size"] = r.status_code, len(data)
    ctype = r.headers.get("Content-Type", "").split(";")[0]
    expected = r.headers.get("Content-Length")

    if r.status_code == 404:
        return {**res, "status": "BROKEN", "problem": "이미지 누락 (404, 파일 없음)"}
    if r.status_code >= 400:
        return {**res, "status": "BROKEN", "problem": f"이미지 열기 실패 (HTTP {r.status_code})"}
    if len(data) == 0:
        return {**res, "status": "BROKEN", "problem": "빈 파일 (0바이트)"}
    if expected and expected.isdigit() and not r.headers.get("Content-Encoding") and int(expected) != len(data):
        return {**res, "status": "BROKEN", "problem": f"전송 중 잘림 ({len(data)}/{expected} bytes)"}
    head = data[:256].lstrip().lower()
    if head.startswith((b"<!doctype", b"<html", b"<?xml", b"{")) and "svg" not in ctype:
        return {**res, "status": "BROKEN", "problem": "이미지 대신 오류 페이지 응답 (누락)"}
    if "svg" in ctype:
        return res

    try:
        with Image.open(io.BytesIO(data)) as im:
            im.verify()
        with Image.open(io.BytesIO(data)) as im:
            res["dim"] = f"{im.size[0]}x{im.size[1]}"
            im.load()
            if getattr(im, "n_frames", 1) > 1:
                im.seek(im.n_frames - 1)
                im.load()
                im.seek(0)
            if min(im.size) <= 2:
                return {**res, "status": "WARN", "problem": f"비정상적으로 작은 이미지 ({res['dim']})"}
            if gray_band_bottom(im):
                return {**res, "status": "WARN", "problem": "하단 회색 띠 - 잘린 이미지 의심 (육안 확인)"}
    except Image.DecompressionBombError:
        return {**res, "status": "WARN", "problem": "초대형 이미지 (검사 생략)"}
    except (OSError, SyntaxError, ValueError) as e:
        msg = str(e).lower()
        if "truncated" in msg:
            return {**res, "status": "BROKEN", "problem": "파일이 중간에 잘림 (일부만 표시됨)"}
        if "cannot identify" in msg and ("avif" in ctype or url.lower().split("?")[0].endswith(".avif")):
            return {**res, "status": "WARN", "problem": "AVIF 검사 미지원 (pip install pillow-avif-plugin)"}
        if "cannot identify" in msg:
            return {**res, "status": "BROKEN", "problem": "손상된 파일 (이미지로 열리지 않음)"}
        return {**res, "status": "BROKEN", "problem": f"손상된 파일 ({str(e)[:50]})"}
    return res


def find_alternate(url, shop_domain, timeout):
    """깨진 이미지가 다른 도메인에 있을 때, 쇼핑몰 도메인의 같은 경로에 정상 파일이 있는지 확인."""
    p = urlparse(url)
    if p.netloc == shop_domain:
        return ""
    alt = f"https://{shop_domain}{p.path}" + (f"?{p.query}" if p.query else "")
    r = check_image(alt, f"https://{shop_domain}/", timeout, 0)
    return alt if r["status"] == "OK" else ""


# ------------------------------------------------------------------ 엑셀
def write_excel(path, products, img_results, started):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    head_fill = PatternFill("solid", fgColor="1F2937")
    head_font = Font(bold=True, color="FFFFFF")
    red = PatternFill("solid", fgColor="FDE2E1")
    yel = PatternFill("solid", fgColor="FFF4D6")
    link_font = Font(color="1D4ED8", underline="single")

    def header(ws, cols, widths):
        ws.append(cols)
        for i, w in enumerate(widths, 1):
            c = ws.cell(row=1, column=i)
            c.fill, c.font = head_fill, head_font
            c.alignment = Alignment(vertical="center")
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}1"

    # 시트 1: 문제 이미지
    ws = wb.active
    ws.title = "문제 이미지"
    header(ws, ["구분", "상품명", "상품번호", "상품 링크", "문제 유형", "이미지 위치", "이미지 URL", "HTTP", "정상 파일 위치(교체용)"],
           [8, 40, 10, 48, 36, 11, 70, 7, 60])
    rows = []
    for no, p in products.items():
        if p["error"]:
            rows.append(("깨짐", p, "", p["error"], "", "", ""))
        for note in p["notes"]:
            rows.append(("주의", p, "", note, "", "", ""))
        for u, where in p["images"].items():
            r = img_results.get(u)
            if r and r["status"] != "OK":
                rows.append(("깨짐" if r["status"] == "BROKEN" else "주의", p, u, r["problem"], where, r["http"], r.get("alt", "")))
    rows.sort(key=lambda x: (x[0] != "깨짐", int(x[1]["no"])))
    for kind, p, u, prob, where, http, alt in rows:
        ws.append([kind, p["name"], int(p["no"]), p["url"], prob, where, u, http, alt])
        rr = ws.max_row
        ws.cell(rr, 4).hyperlink, ws.cell(rr, 4).font = p["url"], link_font
        if u:
            ws.cell(rr, 7).hyperlink, ws.cell(rr, 7).font = u, link_font
        ws.cell(rr, 1).fill = red if kind == "깨짐" else yel
    if not rows:
        ws.append(["", "문제 있는 이미지가 없습니다."])

    # 시트 2: 상품별 요약
    ws2 = wb.create_sheet("상품별 요약")
    header(ws2, ["상품번호", "상품명", "상품 링크", "검사 이미지 수", "깨짐", "주의", "결과"],
           [10, 40, 48, 13, 8, 8, 10])
    for no in sorted(products, key=int):
        p = products[no]
        st = [img_results.get(u, {}).get("status") for u in p["images"]]
        b = st.count("BROKEN") + (1 if p["error"] else 0)
        w = st.count("WARN") + len(p["notes"])
        ws2.append([int(no), p["name"], p["url"], len(p["images"]), b, w,
                    "깨짐" if b else ("주의" if w else "정상")])
        rr = ws2.max_row
        ws2.cell(rr, 3).hyperlink, ws2.cell(rr, 3).font = p["url"], link_font
        if b:
            ws2.cell(rr, 7).fill = red
        elif w:
            ws2.cell(rr, 7).fill = yel

    # 시트 3: 검사 정보
    ws3 = wb.create_sheet("검사 정보")
    broken_products = sum(1 for p in products.values()
                          if p["error"] or any(img_results.get(u, {}).get("status") == "BROKEN" for u in p["images"]))
    for k, v in [("검사 시작", started.strftime("%Y-%m-%d %H:%M")),
                 ("소요 시간(분)", round((datetime.now() - started).seconds / 60, 1)),
                 ("검사 상품 수", len(products)),
                 ("검사 이미지 수(고유)", len(img_results)),
                 ("깨진 이미지가 있는 상품 수", broken_products),
                 ("깨진 이미지 수", sum(1 for r in img_results.values() if r["status"] == "BROKEN")),
                 ("주의 이미지 수", sum(1 for r in img_results.values() if r["status"] == "WARN"))]:
        ws3.append([k, v])
    ws3.column_dimensions["A"].width, ws3.column_dimensions["B"].width = 28, 22
    wb.save(path)
    return broken_products


# ------------------------------------------------------------------ main
def main():
    ap = argparse.ArgumentParser(description="카페24 쇼핑몰 상품 이미지 깨짐 검사")
    ap.add_argument("url", nargs="?", default="https://vivacious.kr/")
    ap.add_argument("--limit", type=int, default=0, help="검사할 최대 상품 수 (0=전체)")
    ap.add_argument("--workers", type=int, default=6, help="동시 요청 수 (기본 6)")
    ap.add_argument("--max-list-pages", type=int, default=3000, help="탐색할 최대 목록 페이지 수")
    ap.add_argument("--delay", type=float, default=0.2, help="목록 페이지 요청 간 대기(초)")
    ap.add_argument("--timeout", type=float, default=25)
    ap.add_argument("--retries", type=int, default=1)
    ap.add_argument("--out", default="")
    args = ap.parse_args()

    base = args.url if args.url.startswith("http") else "https://" + args.url
    domain = urlparse(base).netloc
    started = datetime.now()
    out = args.out or f"{domain.split('.')[0]}_이미지검사_{started:%Y%m%d_%H%M}.xlsx"
    if not out.lower().endswith(".xlsx"):
        out += ".xlsx"

    single_url, single_no = canon_product(base)
    if single_url:  # 상품 링크를 넣으면 그 상품만 검사
        print(f"[1/3] 상품 1개 검사: {single_url}")
        found = {single_no: single_url}
    else:
        print(f"[1/3] 상품 목록 수집: {base}")
        found = discover_products(base, args.workers, args.max_list_pages, args.delay)
    nos = sorted(found, key=int)
    if args.limit:
        nos = nos[:args.limit]
    print(f"  → 상품 {len(nos)}개 검사 예정\n")
    if not nos:
        print("상품을 찾지 못했습니다. 주소를 확인해 주세요.")
        sys.exit(2)

    print("[2/3] 상품 페이지 분석")
    products = {}
    with ThreadPoolExecutor(max_workers=args.workers) as ex:
        futs = {ex.submit(parse_product, found[no], domain): no for no in nos}
        for i, f in enumerate(as_completed(futs), 1):
            no = futs[f]
            p = f.result()
            p["no"] = no
            products[no] = p
            if i % 25 == 0 or i == len(nos):
                print(f"  {i}/{len(nos)}")

    all_imgs = {}
    for p in products.values():
        for u in p["images"]:
            all_imgs.setdefault(u, p["url"])
    print(f"  → 고유 이미지 {len(all_imgs)}개\n")

    print("[3/3] 이미지 검사")
    img_results = {}
    with ThreadPoolExecutor(max_workers=args.workers) as ex:
        futs = {ex.submit(check_image, u, ref, args.timeout, args.retries): u for u, ref in all_imgs.items()}
        for i, f in enumerate(as_completed(futs), 1):
            u = futs[f]
            img_results[u] = f.result()
            if img_results[u]["status"] == "BROKEN":
                print(f"  ✕ {img_results[u]['problem']}  {u}")
            if i % 200 == 0 or i == len(all_imgs):
                print(f"  {i}/{len(all_imgs)}")

    broken = [u for u, r in img_results.items() if r["status"] == "BROKEN" and urlparse(u).netloc != domain]
    if broken:
        print(f"  다른 도메인의 깨진 이미지 {len(broken)}개 → 쇼핑몰 주소에 같은 파일이 있는지 확인 중")
        with ThreadPoolExecutor(max_workers=args.workers) as ex:
            for u, alt in zip(broken, ex.map(lambda x: find_alternate(x, domain, args.timeout), broken)):
                img_results[u]["alt"] = alt

    if single_url:
        p = products[single_no]
        print(f"\n상품명: {p['name']}")
        if p["error"]:
            print(f"  [깨짐] {p['error']}")
        for note in p["notes"]:
            print(f"  [주의] {note}")
        for u, where in p["images"].items():
            r = img_results[u]
            label = {"OK": "정상", "BROKEN": "깨짐", "WARN": "주의"}[r["status"]]
            print(f"  [{label}] {where:5} {r['dim']:>10}  {r['problem']}  {u}")
            if r.get("alt"):
                print(f"           └ 정상 파일 있음: {r['alt']}")

    bp = write_excel(out, products, img_results, started)
    nb = sum(1 for r in img_results.values() if r["status"] == "BROKEN")
    nw = sum(1 for r in img_results.values() if r["status"] == "WARN")
    print(f"\n완료: 상품 {len(products)}개 / 이미지 {len(img_results)}개 검사")
    print(f"  깨진 이미지 {nb}개 (상품 {bp}개), 주의 {nw}개")
    print(f"  엑셀: {out}")


if __name__ == "__main__":
    main()
